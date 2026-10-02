"""Excel (openpyxl).

Summary: ONE sheet "Purchase orders": the header lines in the first rows, a blank row, the table (bold header, frozen below it,
         autofilter).
Detail:  a first sheet "About" with the header lines, then one sheet per section (key/value sections as Field / Value).
Money cells are numbers written from exact Decimal with the format #,##0.00; quantities are numbers; text cells go through
`safe_cell` and are forced to the string type, so a value starting with = is never stored as a formula.
"""
import io
import re
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.po.export.model import ExportDoc, KeyValues, Table
from app.po.export.safety import safe_cell

MONEY_FORMAT = "#,##0.00"
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")
BOLD = Font(bold=True)


def _put(ws, row: int, col: int, value: Any, kind: str) -> None:
    cell = ws.cell(row=row, column=col)
    if value is None:
        return
    if kind == "money":
        cell.value = Decimal(value)
        cell.number_format = MONEY_FORMAT
    elif kind in ("qty", "price"):
        cell.value = Decimal(value)
    elif kind == "int":
        cell.value = int(value)
    else:
        cell.value = safe_cell(str(value))
        cell.data_type = "s"                                              # never a formula, whatever the text
        if len(cell.value) > 60:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(width + 2, 60))


def _meta(ws, doc: ExportDoc) -> int:
    for i, (label, value) in enumerate(doc.meta, start=1):
        _put(ws, i, 1, label, "text")
        ws.cell(row=i, column=1).font = BOLD
        _put(ws, i, 2, value, "text")
    return len(doc.meta)


def _table(ws, section: Table, top: int, autofilter: bool) -> None:
    widths: dict[int, int] = {}
    for j, col in enumerate(section.columns, start=1):
        _put(ws, top, j, col.label, "text")
        ws.cell(row=top, column=j).font = BOLD
        widths[j] = len(col.label)
    if not section.rows:
        _put(ws, top + 1, 1, "None", "text")
    for i, row in enumerate(section.rows, start=top + 1):
        for j, (value, col) in enumerate(zip(row, section.columns), start=1):
            _put(ws, i, j, value, col.kind)
            widths[j] = max(widths[j], len(str(value)) if value is not None else 0)
    if section.note:
        _put(ws, top + 1 + max(len(section.rows), 1), 1, section.note, "text")
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    if autofilter and section.rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(section.columns))}{top + len(section.rows)}"
    _widths(ws, widths)


def _kv(ws, section: KeyValues) -> None:
    _put(ws, 1, 1, "Field", "text")
    _put(ws, 1, 2, "Value", "text")
    ws.cell(row=1, column=1).font = ws.cell(row=1, column=2).font = BOLD
    for i, (label, value, kind) in enumerate(section.items, start=2):
        _put(ws, i, 1, label, "text")
        _put(ws, i, 2, value, kind)
    _widths(ws, {1: max([len(x[0]) for x in section.items] + [5]), 2: 40})


def sheet_name(title: str, used: set[str]) -> str:
    base = _BAD_SHEET.sub("", title)[:31] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:28]} {n}"
        n += 1
    used.add(name.lower())
    return name


SHORT_TITLES = {"Invoices matched to this PO": "Invoices", "How the commits are allocated": "Allocations",
                "Where this PO came from": "Provenance"}


def render(doc: ExportDoc) -> bytes:
    wb = Workbook()
    ws = wb.active
    if doc.kind == "summary":
        ws.title = "Purchase orders"
        top = _meta(ws, doc) + 2
        _table(ws, doc.sections[0], top, autofilter=True)
    else:
        used: set[str] = set()
        ws.title = sheet_name("About", used)
        _meta(ws, doc)
        _widths(ws, {1: 18, 2: 50})
        for section in doc.sections:
            sheet = wb.create_sheet(sheet_name(SHORT_TITLES.get(section.title, section.title), used))
            if isinstance(section, KeyValues):
                _kv(sheet, section)
            else:
                _table(sheet, section, 1, autofilter=False)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
