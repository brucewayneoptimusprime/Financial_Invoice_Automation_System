"""PO export, stage E2: PDF and Word, plus the full matrix (summary / financial / full x pdf / docx / xlsx / csv): every combination
answers 200 with the right type, opens in its reader, and writes nothing. PDF and Word numbers are compared with the screens' JSON."""
import csv
import io
from decimal import Decimal

import docx
import openpyxl
import pypdfium2 as pdfium
import pytest

from tests.api.helpers import api
from tests.gmail.helpers import table_counts
from tests.po.export_helpers import PO_SS_001, approved_api, insert_po

MEDIA = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8"}
ENTRIES = {"summary": ("/api/pos/export", {}), "financial": (f"/api/pos/{PO_SS_001}/export", {"level": "financial"}),
           "full": (f"/api/pos/{PO_SS_001}/export", {"level": "full"})}


def fmt(value) -> str:
    return f"{Decimal(str(value)):,.2f}"


def pdf_text(body: bytes) -> tuple[str, int]:
    pdf = pdfium.PdfDocument(body)
    text = "\n".join(pdf[i].get_textpage().get_text_bounded() for i in range(len(pdf)))
    return text, len(pdf)


def opens(fmt_: str, body: bytes) -> bool:
    if fmt_ == "pdf":
        return len(pdfium.PdfDocument(body)) >= 1
    if fmt_ == "docx":
        return len(docx.Document(io.BytesIO(body)).paragraphs) > 0
    if fmt_ == "xlsx":
        return len(openpyxl.load_workbook(io.BytesIO(body)).sheetnames) >= 1
    return len(list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))) > 3


# ------------------------------------------------------------------------------------------ the matrix

@pytest.fixture(scope="module")
def matrix(tmp_path_factory):
    """Every entry x format exported once from one approved database, with the table counts before and after."""
    tmp = tmp_path_factory.mktemp("matrix")
    with pytest.MonkeyPatch.context() as mp:                     # module scope: blank the secrets like conftest, restored afterwards
        for k in ("ANTHROPIC_API_KEY", "GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET", "OAUTH_ENCRYPTION_KEY", "GMAIL_BACKEND"):
            mp.setenv(k, "")
        from app.config import get_settings
        get_settings.cache_clear()
        with approved_api(tmp) as c:
            before = table_counts(c.db_path)
            out = {(e, f): c.get(path, params={"format": f, **params}) for e, (path, params) in ENTRIES.items() for f in MEDIA}
            after = table_counts(c.db_path)
            detail, listing = c.get(f"/api/pos/{PO_SS_001}").json(), c.get("/api/pos").json()["pos"]
        get_settings.cache_clear()
    return out, before, after, detail, listing


@pytest.mark.parametrize("entry", list(ENTRIES))
@pytest.mark.parametrize("format_", list(MEDIA))
def test_every_entry_level_and_format_is_built_and_opens(matrix, entry, format_):
    out, *_ = matrix
    r = out[(entry, format_)]
    assert r.status_code == 200 and r.headers["content-type"] == MEDIA[format_]
    assert r.headers["content-disposition"].startswith("attachment;") and r.headers["content-disposition"].split('"')[1].endswith(f".{format_}")
    assert opens(format_, r.content)


def test_twelve_exports_wrote_nothing(matrix):
    _, before, after, *_ = matrix
    assert before == after


# ------------------------------------------------------------------------------------------ PDF: values equal the screen

def test_the_pdf_shows_the_po_page_values(matrix):
    out, _, _, d, _ = matrix
    text, pages = pdf_text(out[("full", "pdf")].content)
    a = d["amounts"]
    for value in (a["total"], a["committed"], a["balance"], a["awaiting_review"], a["consumed_without_line"]):
        assert fmt(value) in text
    for ln in d["lines"]:
        assert fmt(ln["amount"]) in text and fmt(ln["consumed_amount"]) in text and str(ln["quantity"]) in text
    for e in d["ledger"]:
        assert fmt(e["amount"]) in text
    for x in d["allocations"]:
        assert fmt(x["amount"]) in text
    assert "PO-SS-001" in text and "How the commits are allocated" in text and "Where this PO came from" in text
    assert f"Page 1 of {pages}" in text and "Amount (USD)" in text and "Unit price" in text   # a narrow header may wrap its "(USD)"


def test_the_financial_pdf_has_no_metadata_sections(matrix):
    out, *_ = matrix
    text, _ = pdf_text(out[("financial", "pdf")].content)
    assert "Ledger" in text and "How the commits are allocated" not in text and "Where this PO came from" not in text


def test_the_summary_pdf_lists_every_po_with_its_amounts_and_the_scope(matrix):
    out, *_, listing = matrix
    text, _ = pdf_text(out[("summary", "pdf")].content)
    assert "Scope: All purchase orders" in text and "Nothing was sent anywhere" in text
    for p in listing:
        assert p["po_number"] in text and fmt(p["total"]) in text and fmt(p["balance"]) in text


# ------------------------------------------------------------------------------------------ Word: values equal the screen

def _docx_tables(body: bytes):
    d = docx.Document(io.BytesIO(body))
    return d, [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]


def _by_header(tables, first: str):
    """The table whose header row starts with `first` (empty sections are a "None" paragraph, not a table)."""
    return next(t for t in tables if t and t[0] and t[0][0] == first)


def test_the_word_file_shows_the_po_page_values_with_repeating_headers(matrix):
    out, _, _, d, _ = matrix
    document, tables = _docx_tables(out[("full", "docx")].content)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Purchase order PO-SS-001", "PO", "Totals", "Invoices matched to this PO", "Lines", "Ledger",
                        "How the commits are allocated", "Where this PO came from"]
    totals = dict((row[0], row[1]) for row in tables[1])
    assert totals["Balance"] == f"{fmt(d['amounts']['balance'])} USD" and totals["Total"] == f"{fmt(d['amounts']['total'])} USD"
    lines = tables[3]
    assert lines[0][:5] == ["Line", "Description", "Quantity", "Unit price (USD)", "Amount (USD)"]
    assert lines[1][4] == fmt(d["lines"][0]["amount"]) and lines[1][1] == d["lines"][0]["description"]
    assert [row[1] for row in tables[4][1:]] == [fmt(e["amount"]) for e in d["ledger"]]
    header_rows = [t.rows[0] for t in document.tables[2:5]]
    assert all(r._tr.trPr is not None and r._tr.trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is not None
               for r in header_rows)


def test_the_summary_word_file_is_one_table(matrix):
    out, *_, listing = matrix
    document, tables = _docx_tables(out[("summary", "docx")].content)
    assert len(tables) == 1 and tables[0][0][:4] == ["PO number", "Vendor", "Currency", "Total"]
    assert [row[0] for row in tables[0][1:]] == [p["po_number"] for p in listing]
    assert [row[3] for row in tables[0][1:]] == [fmt(p["total"]) for p in listing]
    assert any(p.text == "Scope: All purchase orders" for p in document.paragraphs)


# ------------------------------------------------------------------------------------------ text the PDF font cannot draw; long POs

def test_characters_helvetica_cannot_draw_become_question_marks_with_a_note_and_word_keeps_them(tmp_path):
    with api(tmp_path) as c:
        pid = insert_po(c, po_number="PO-UNI-1", vendor="Ünïcode Trading ₹ 漢字", lines=[("Café crème \x07 ✓ widgets", "1", "2.00", 200)])
        pdf = c.get(f"/api/pos/{pid}/export", params={"format": "pdf", "level": "full"}).content
        word = c.get(f"/api/pos/{pid}/export", params={"format": "docx", "level": "full"}).content
        plain = c.get(f"/api/pos/{pid}/export", params={"format": "pdf", "level": "full"}).content
    text, _ = pdf_text(pdf)
    assert "Ünïcode Trading ? ?? (approved)" in text and "Café crème ? widgets" in text     # PDF text collapses double spaces
    assert 'shown as "?"' in text
    document, tables = _docx_tables(word)
    assert "Ünïcode Trading ₹ 漢字 (approved)" in [row[1] for row in tables[0]]
    assert _by_header(tables, "Line")[1][1] == "Café crème  ✓ widgets"                 # the control character dropped, the rest kept
    assert plain


def test_a_pdf_without_unusual_characters_has_no_replacement_note(matrix):
    out, *_ = matrix
    text, _ = pdf_text(out[("full", "pdf")].content)
    assert 'shown as "?"' not in text


def test_a_very_long_po_paginates_with_repeated_headers(tmp_path):
    long_lines = [(f"Item {n} " + "very long description " * 45, "3", "12.34", 3702) for n in range(200)]
    with api(tmp_path) as c:
        pid = insert_po(c, po_number="PO-LONG", vendor="Long Supplies", lines=long_lines, total=200 * 3702)
        pdf = c.get(f"/api/pos/{pid}/export", params={"format": "pdf", "level": "full"}).content
        word = c.get(f"/api/pos/{pid}/export", params={"format": "docx", "level": "full"}).content
    doc = pdfium.PdfDocument(pdf)
    assert len(doc) > 10
    page_texts = [doc[i].get_textpage().get_text_bounded() for i in range(len(doc))]
    assert sum("Description" in t and "Remaining" in t for t in page_texts[1:]) > 5           # the Lines header on following pages
    assert all(f"of {len(doc)}" in t for t in page_texts)
    _, tables = _docx_tables(word)
    assert len(_by_header(tables, "Line")) == 201


def test_an_empty_po_in_pdf_and_word(tmp_path):
    with api(tmp_path) as c:
        pid = insert_po(c, po_number="PO-EMPTY-2", vendor="Nothing Yet")
        pdf = c.get(f"/api/pos/{pid}/export", params={"format": "pdf", "level": "full"}).content
        word = c.get(f"/api/pos/{pid}/export", params={"format": "docx", "level": "full"}).content
    text, _ = pdf_text(pdf)
    assert text.count("None") >= 4
    assert sum(p.text == "None" for p in docx.Document(io.BytesIO(word)).paragraphs) == 4
