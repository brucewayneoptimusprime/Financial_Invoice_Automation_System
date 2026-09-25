"""PO document readers (PO integration stage 4). All documents are generated here; no real client documents."""
import io
import zipfile

import openpyxl
import pytest

from app.po.readers import OLE_MAGIC, PODocRejected, read_po_document
from tests.api.helpers import api_settings
from tests.ingest.docs import make_native_pdf

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DRAFT = "a" * 32


def docx_bytes(paragraphs, table=None, extra=None) -> bytes:
    def p(text):
        return f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
    body = "".join(p(t) for t in paragraphs)
    if table:
        rows = "".join("<w:tr>" + "".join(f"<w:tc>{p(c)}</w:tc>" for c in row) + "</w:tr>" for row in table)
        body += f"<w:tbl>{rows}</w:tbl>"
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", xml)
        for name, data in (extra or {}).items():
            z.writestr(name, data)
    return buf.getvalue()


def write(tmp_path, name, data: bytes):
    p = tmp_path / name
    p.write_bytes(data)
    return p


def read(tmp_path, name, data, **settings_kw):
    return read_po_document(write(tmp_path, name, data), name, DRAFT, api_settings(tmp_path, **settings_kw))


def test_docx_paragraphs_and_table_rows_in_order(tmp_path):
    doc = read(tmp_path, "po.docx", docx_bytes(["PURCHASE ORDER PO-1001", "Supplier: Acme Ltd"],
                                               [["Item", "Qty", "Price"], ["Widget", "10", "5.00"]]))
    assert doc.kind == "text" and doc.text_usable and doc.pages[0].image_data is None
    assert doc.page_texts[1] == "PURCHASE ORDER PO-1001\nSupplier: Acme Ltd\nItem | Qty | Price\nWidget | 10 | 5.00"
    assert (tmp_path / "po_drafts" / DRAFT / "document.txt").is_file() and len(doc.sha256) == 64


def test_xlsx_reads_cached_values_and_never_evaluates_a_formula(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO"
    ws.append(["PO number", "PO-2002"])
    ws.append(["Qty", 3, "Price", 7])
    ws["E2"] = "=B2*D2"                                                 # no cached value: openpyxl never computed it
    buf = io.BytesIO()
    wb.save(buf)
    doc = read(tmp_path, "po.xlsx", buf.getvalue())
    text = doc.page_texts[1]
    assert "Sheet: PO" in text and "PO number | PO-2002" in text and "21" not in text and "=B2*D2" not in text


def test_xlsx_is_cut_at_the_cell_cap(tmp_path):
    wb = openpyxl.Workbook()
    for i in range(50):
        wb.active.append([f"r{i}", i, i * 2])
    buf = io.BytesIO()
    wb.save(buf)
    doc = read(tmp_path, "big.xlsx", buf.getvalue(), po_sheet_max_cells=30)
    assert any("cut at 30 cells" in n for n in doc.notes) and "r20" not in doc.page_texts[1]


def test_csv_with_a_sniffed_delimiter(tmp_path):
    doc = read(tmp_path, "po.csv", "PO number;PO-3003\nItem;Qty;Price\nChair;2;99.50\n".encode("utf-8"))
    assert doc.page_texts[1] == "PO number | PO-3003\nItem | Qty | Price\nChair | 2 | 99.50"


def test_a_pdf_goes_through_the_invoice_ingest(tmp_path):
    pdf = make_native_pdf(tmp_path / "po.pdf", [["PURCHASE ORDER", "PO Number: PO-4004", "Supplier: Acme Ltd", "Total: 1,250.00 USD"]])
    doc = read_po_document(pdf, "po.pdf", DRAFT, api_settings(tmp_path))
    assert doc.kind == "rendered" and doc.media_type == "application/pdf" and doc.text_usable and doc.path == "text_and_vision"
    assert doc.pages[0].image_data and "PO-4004" in doc.page_texts[1]
    assert (tmp_path / "po_drafts" / DRAFT / "pages").is_dir()


def test_a_password_protected_pdf_is_unreadable_without_a_model_call(tmp_path):
    pdf = make_native_pdf(tmp_path / "locked.pdf", [["PO-5005"]], password="pw")
    doc = read_po_document(pdf, "locked.pdf", DRAFT, api_settings(tmp_path))
    assert doc.failure_code == "password_protected" and doc.pages == []


@pytest.mark.parametrize("name, data, code", [
    ("old.doc", OLE_MAGIC + b"\0" * 600, "legacy_office"),
    ("old.xls", OLE_MAGIC + b"\0" * 600, "legacy_office"),
    ("macro.docm", docx_bytes(["x"], extra={"word/vbaProject.bin": b"\0" * 10}), "macro_enabled"),
    ("other.zip", b"PK\x03\x04" + b"\0" * 10, "unsupported_type"),
    ("notes.txt", b"just some text", "unsupported_type"),
    ("empty.csv", b"", "empty_file"),
    ("blank.docx", docx_bytes([]), "blank_document"),
    ("bin.csv", b"a,b\x00c", "unsupported_type"),
])
def test_refused_documents(tmp_path, name, data, code):
    with pytest.raises(PODocRejected) as exc:
        read(tmp_path, name, data)
    assert exc.value.code == code and exc.value.message


def test_a_zip_bomb_is_refused_before_it_is_expanded(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "a" * 5_000_000)
    with pytest.raises(PODocRejected) as exc:
        read(tmp_path, "bomb.docx", buf.getvalue(), po_zip_max_uncompressed_bytes=1_000_000)
    assert exc.value.code == "too_large"


def test_long_text_is_cut_with_a_note(tmp_path):
    doc = read(tmp_path, "long.csv", ("x," * 30000).encode(), po_doc_text_max_chars=1000)
    assert len(doc.page_texts[1]) == 1000 and any("cut at 1000" in n for n in doc.notes)
